"""
Modo Marketplace — precificação para Shopee e Mercado Livre.

Módulo independente do fluxo varejo/atacado: reaproveita apenas a leitura
de XML/CSV (`gerar_tabela`) de processador.py, sem alterar nada lá.

As vendas online saem por CNPJ do Simples Nacional: um único % (DAS) sobre
o preço de venda, sem Federal/ICMS separados e sem crédito de ICMS na entrada.

Fluxo por produto:
    1. CUSTO ENTRADA  — NF × multiplicador + ST + ANT + IPI + frete (% do custo real)
    2. CUSTO BASE     — C. ENTRADA + embalagem de envio
    3. PREÇO BASE     — preço que atinge a margem alvo só com Simples + despesa
                        (antes das taxas de marketplace)
    4. PREÇO SHOPEE / PREÇO ML — preço que atinge a margem alvo pagando as
                        taxas da faixa em que o próprio preço cai
Simples e despesa operacional são % do preço de venda de cada canal.
"""
import html
import math

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── Tabelas de taxas (anotação 2024 — conferir periodicamente) ──────────────
# (a partir de R$, tarifa fixa R$, comissão %)
SHOPEE_FAIXAS = [
    (0.00,   4.00, 0.20),
    (80.00, 16.00, 0.14),
    (100.00, 20.00, 0.14),
    (200.00, 26.00, 0.14),   # anotação vai até R$ 499 — acima disso assume a mesma faixa
]

# Mercado Livre Premium: comissão (parâmetro, padrão 16,5%) + faixa abaixo.
# (a partir de R$, tarifa fixa R$, % extra sobre o preço)
ML_FAIXAS = [
    (0.00,  0.00, 0.50),     # até R$ 12,50: + 50% do valor do item
    (12.51, 6.25, 0.00),
    (29.01, 6.75, 0.00),
    (50.01, 6.75, 0.00),
    (79.00, 0.00, 0.00),     # a partir de R$ 79: sem tarifa fixa, vendedor paga o frete
]
ML_LIMITE_FRETE = 79.00

# Desconto no frete do ML por reputação do vendedor
REPUTACAO_DESC = {'verde': 0.50, 'amarela': 0.40, 'vermelha': 0.0}

MELHOR_SHOPEE = 'SHOPEE'
MELHOR_ML     = 'MERCADO LIVRE'


# ─── Parâmetros ──────────────────────────────────────────────────────────────
def montar_params(form, frete_pct=None):
    """
    Converte o formulário do modo marketplace em parâmetros numéricos.
    `frete_pct` sobrescreve o % de frete (gerar_tabela zera em NF CIF).
    """
    def num(key, default):
        try:
            return float(str(form.get(key, default)).replace(',', '.'))
        except (TypeError, ValueError):
            return float(default)

    reputacao = str(form.get('reputacao_ml', 'verde')).lower()
    if reputacao not in REPUTACAO_DESC:
        reputacao = 'verde'

    return {
        'mult':      num('mult_mkt', 2.0),
        'frete':     frete_pct if frete_pct is not None else num('frete', 10) / 100,
        'desp':      num('desp_mkt', 5) / 100,
        'simples':   num('simples', 15) / 100,
        'margem':    num('margem_mkt', 15) / 100,
        'emb':       num('emb_mkt', 0.0),
        'com_ml':    num('com_ml', 16.5) / 100,
        'frete_ml':  num('frete_ml', 20.0),
        'reputacao': reputacao,
    }


def frete_ml_item(M):
    """Frete que o vendedor paga no ML (preço ≥ R$ 79) já com desconto de reputação."""
    return round(M['frete_ml'] * (1 - REPUTACAO_DESC[M['reputacao']]), 2)


# ─── Custos ──────────────────────────────────────────────────────────────────
def custo_entrada(row, M):
    """
    C. ENTRADA do marketplace. Simples Nacional não aproveita crédito de ICMS;
    ST e ANT já pagos na entrada entram como custo. Frete sobre o custo real.
    """
    c_real = round(row['nf_u'] * M['mult'], 2)
    frete  = round(c_real * M['frete'], 2)
    c_ent  = round(c_real + row['st_u'] + row['ant_u'] + row['ipi_u'] + frete, 2)
    return dict(c_real=c_real, frete=frete, c_ent=c_ent)


def deducoes_venda(preco, M):
    """(simples, despesa) da venda, cada um arredondado — igual às colunas do Excel."""
    return round(preco * M['simples'], 2), round(preco * M['desp'], 2)


def _faixa(faixas, preco):
    atual = faixas[0]
    for f in faixas:
        if preco >= f[0]:
            atual = f
    return atual


def faixas_ml(M):
    """
    Faixas do ML já com a comissão somada ao % da faixa e a tarifa fixa da
    última faixa trocada pelo frete do vendedor.
    """
    frete = frete_ml_item(M)
    return [(lo, frete if lo >= ML_LIMITE_FRETE else fixo, M['com_ml'] + pct)
            for lo, fixo, pct in ML_FAIXAS]


def taxas_shopee(preco):
    _, fixo, pct = _faixa(SHOPEE_FAIXAS, preco)
    return round(preco * pct + fixo, 2)


def taxas_ml(preco, M):
    _, fixo, pct = _faixa(faixas_ml(M), preco)
    return round(preco * pct + fixo, 2)


# ─── Resolução de preço ──────────────────────────────────────────────────────
def _resolver(custo, fixo, pct, M):
    """
    Preço P tal que  P − custo − fixo − (pct + simples + desp)·P = margem·P.
    Retorna None se a margem é inatingível (denominador ≤ 0).
    """
    den = 1 - pct - M['simples'] - M['desp'] - M['margem']
    return (custo + fixo) / den if den > 0 else None


def _preco_por_faixas(custo, faixas, M):
    """Acha a faixa em que o próprio preço cai (as taxas dependem do preço)."""
    for i, (lo, fixo, pct) in enumerate(faixas):
        hi = faixas[i + 1][0] if i + 1 < len(faixas) else math.inf
        p = _resolver(custo, fixo, pct, M)
        if p is None:
            continue
        if p < lo:
            # Faixa anterior não fechava a conta e nesta sobra margem → piso da faixa
            return lo
        if p < hi:
            return p
    return None


def arredondar_x9_acima(val):
    """Arredonda PARA CIMA até o próximo X,X9 (nunca fica abaixo do preço mínimo)."""
    if val <= 0:
        return 0.0
    return round(math.ceil(round((val - 0.09) * 10, 6)) / 10 + 0.09, 2)


def _arredondar_na_faixa(preco, faixas):
    """X,X9 acima, sem pular para a faixa seguinte (que teria taxa maior)."""
    arred = arredondar_x9_acima(preco)
    if _faixa(faixas, arred) is not _faixa(faixas, preco):
        arred = math.ceil(round(preco * 100, 6)) / 100
    return round(arred, 2)


_CANAL_VAZIO = dict(preco=0.0, com_pct=0.0, comissao=0.0, tarifa=0.0, frete=0.0,
                   simples=0.0, desp=0.0, taxas=0.0, lucro=0.0, margem=0.0)


def _canal(preco, custo_base, faixas, M, limite_frete=None):
    """
    Detalha a venda num canal: comissão (% da faixa), tarifa fixa, frete do
    vendedor (ML ≥ R$ 79 — ocupa o lugar da tarifa fixa na faixa), Simples e despesa.
    """
    _, fixo, pct = _faixa(faixas, preco)
    frete = fixo if limite_frete is not None and preco >= limite_frete else 0.0
    tarifa = 0.0 if frete else fixo
    comissao = round(preco * pct, 2)
    simples, desp = deducoes_venda(preco, M)
    taxas = round(comissao + tarifa + frete, 2)
    lucro = round(preco - custo_base - taxas - simples - desp, 2)
    return dict(preco=preco, com_pct=pct, comissao=comissao, tarifa=tarifa, frete=frete,
                simples=simples, desp=desp, taxas=taxas, lucro=lucro,
                margem=round(lucro / preco, 4) if preco > 0 else 0.0)


def calcular_marketplace(row, M):
    """Todas as métricas de um produto no modo marketplace."""
    ent = custo_entrada(row, M)
    custo_base = round(ent['c_ent'] + M['emb'], 2)

    p_base_raw = _resolver(custo_base, 0.0, 0.0, M)
    p_base = arredondar_x9_acima(p_base_raw) if p_base_raw else 0.0
    base = _canal(p_base, custo_base, [(0.0, 0.0, 0.0)], M) if p_base else dict(_CANAL_VAZIO)

    sh_raw = _preco_por_faixas(custo_base, SHOPEE_FAIXAS, M)
    if sh_raw is None:
        shopee = dict(_CANAL_VAZIO)
    else:
        shopee = _canal(_arredondar_na_faixa(sh_raw, SHOPEE_FAIXAS), custo_base, SHOPEE_FAIXAS, M)

    fx_ml = faixas_ml(M)
    ml_raw = _preco_por_faixas(custo_base, fx_ml, M)
    if ml_raw is None:
        ml = dict(_CANAL_VAZIO)
    else:
        ml = _canal(_arredondar_na_faixa(ml_raw, fx_ml), custo_base, fx_ml, M,
                    limite_frete=ML_LIMITE_FRETE)

    melhor = MELHOR_SHOPEE if shopee['lucro'] >= ml['lucro'] else MELHOR_ML

    return dict(**ent, emb=M['emb'], custo_base=custo_base, p_base=p_base, base=base,
                frete_ml_base=M['frete_ml'],
                margem_alvo=M['margem'],
                shopee=shopee, ml=ml, melhor=melhor,
                inviavel=sh_raw is None or ml_raw is None)


# ─── Serialização (para o passo de seleção no frontend) ─────────────────────
_CAMPOS_ROW = {
    'nf': str, 'desc': str, 'ref': str, 'sku': str, 'cst': str,
    'qtd': float, 'nf_u': float, 'st_u': float, 'ant_u': float, 'ipi_u': float,
    'cred_pct': float, 'qtd_emb': int, 'p_atual': float,
    'tem_st': bool, 'tem_ant': bool, 'rateado': bool,
}


def row_para_json(row):
    """Mantém só os campos usados no marketplace, em tipos nativos (sem numpy)."""
    return {k: conv(row.get(k, conv())) for k, conv in _CAMPOS_ROW.items()}


def row_de_json(dados):
    """Valida/converte uma linha devolvida pelo frontend."""
    return {k: conv(dados.get(k, conv())) for k, conv in _CAMPOS_ROW.items()}


# ─── Colunas (compartilhadas entre prévia HTML e Excel) ─────────────────────
COR_SHOPEE = 'FDE2D0'
COR_ML     = 'FFF3B0'
COR_ENT    = 'E0F2FE'
COR_BASE   = 'EDE9FE'
COR_ST     = 'FCE4D6'
COR_ANT    = 'D1FAE5'
COR_EDIT   = 'E2EFDA'
COR_LINK   = 'F1F5F9'

FMT_BRL  = 'R$ #,##0.00'
FMT_PCT  = '0.00%'
FMT_MULT = '0.00"×"'

# Grupos: (chave, rótulo, cor)
GRUPOS_MKT = [
    ('ID',     '',                               'F8F9FA'),
    ('ENT',    'CUSTO DE ENTRADA (passo a passo)', COR_ENT),
    ('BASE',   'VENDA BASE (sem marketplace)',   COR_BASE),
    ('SHOPEE', 'SHOPEE',                         COR_SHOPEE),
    ('ML',     'MERCADO LIVRE',                  COR_ML),
    ('FIM',    '',                               'F8F9FA'),
]

# Variáveis que ficam no topo da coluna (linha 3 do Excel / 3ª linha do cabeçalho da prévia).
#   ('param', chave de M, formato)      → célula editável; as fórmulas da coluna apontam para ela
#   ('link', coluna de origem)          → espelha a variável de outra coluna (edite na origem)
#   ('texto', rótulo)                   → só informativo (ex.: % vem da tabela de faixas)
# Colunas: (chave, cabeçalho, grupo, formato, variável ou None)
COLUNAS_MKT = [
    ('NF',           'NF',                     'ID',     None,     None),
    ('DESC',         'DESCRIÇÃO',              'ID',     None,     None),
    ('REF',          'REF',                    'ID',     None,     None),
    ('SKU',          'SKU',                    'ID',     None,     None),
    ('QTD',          'QTD',                    'ID',     None,     None),
    ('CST',          'CST',                    'ID',     None,     None),
    ('NF_U',         'NF UNIT',                'ENT',    FMT_BRL,  None),
    ('C_REAL',       '(×) CUSTO REAL',         'ENT',    FMT_BRL,  ('param', 'mult', FMT_MULT)),
    ('ST_U',         '(+) ST',                 'ENT',    FMT_BRL,  None),
    ('ANT_U',        '(+) ANT',                'ENT',    FMT_BRL,  None),
    ('IPI_U',        '(+) IPI',                'ENT',    FMT_BRL,  None),
    ('FRETE',        '(+) FRETE',              'ENT',    FMT_BRL,  ('param', 'frete', FMT_PCT)),
    ('C_ENT',        '(=) CUSTO ENTRADA',      'ENT',    FMT_BRL,  None),
    ('EMB',          '(+) EMBALAGEM',          'ENT',    FMT_BRL,  ('param', 'emb', FMT_BRL)),
    ('C_BASE',       '(=) CUSTO BASE',         'ENT',    FMT_BRL,  None),
    ('MARGEM_ALVO',  'MARGEM ALVO',            'BASE',   FMT_PCT,  ('param', 'margem', FMT_PCT)),
    ('P_BASE',       'PREÇO BASE',             'BASE',   FMT_BRL,  None),
    ('B_SIMPLES',    '(−) SIMPLES',            'BASE',   FMT_BRL,  ('param', 'simples', FMT_PCT)),
    ('B_DESP',       '(−) DESPESA',            'BASE',   FMT_BRL,  ('param', 'desp', FMT_PCT)),
    ('B_LUCRO',      '(=) LUCRO/UN',           'BASE',   FMT_BRL,  None),
    ('SH_PRECO',     'PREÇO SHOPEE',           'SHOPEE', FMT_BRL,  None),
    ('SH_COM_PCT',   'COMISSÃO %',             'SHOPEE', FMT_PCT,  ('texto', 'faixa')),
    ('SH_COM',       '(−) COMISSÃO',           'SHOPEE', FMT_BRL,  None),
    ('SH_TARIFA',    '(−) TARIFA FIXA',        'SHOPEE', FMT_BRL,  ('texto', 'faixa')),
    ('SH_SIMPLES',   '(−) SIMPLES',            'SHOPEE', FMT_BRL,  ('link', 'B_SIMPLES')),
    ('SH_DESP',      '(−) DESPESA',            'SHOPEE', FMT_BRL,  ('link', 'B_DESP')),
    ('SH_LUCRO',     '(=) LUCRO/UN',           'SHOPEE', FMT_BRL,  None),
    ('SH_MARGEM',    'MARGEM',                 'SHOPEE', FMT_PCT,  None),
    ('ML_PRECO',     'PREÇO ML',               'ML',     FMT_BRL,  None),
    ('ML_COM_PCT',   'COMISSÃO %',             'ML',     FMT_PCT,  ('param', 'com_ml', FMT_PCT)),
    ('ML_COM',       '(−) COMISSÃO',           'ML',     FMT_BRL,  None),
    ('ML_TARIFA',    '(−) TARIFA FIXA',        'ML',     FMT_BRL,  ('texto', 'faixa')),
    ('ML_FRETE_BASE', 'FRETE BASE (≥ R$ 79)',  'ML',     FMT_BRL,  ('param', 'frete_ml', FMT_BRL)),
    ('ML_FRETE',     '(−) FRETE ML',           'ML',     FMT_BRL,  ('param', 'desc_rep', FMT_PCT)),
    ('ML_SIMPLES',   '(−) SIMPLES',            'ML',     FMT_BRL,  ('link', 'B_SIMPLES')),
    ('ML_DESP',      '(−) DESPESA',            'ML',     FMT_BRL,  ('link', 'B_DESP')),
    ('ML_LUCRO',     '(=) LUCRO/UN',           'ML',     FMT_BRL,  None),
    ('ML_MARGEM',    'MARGEM',                 'ML',     FMT_PCT,  None),
    ('MELHOR',       'MELHOR CANAL',           'FIM',    None,     None),
]
COL_MKT = {c[0]: i for i, c in enumerate(COLUNAS_MKT, start=1)}
_COR_GRUPO = {g: cor for g, _, cor in GRUPOS_MKT}

# Colunas que o usuário pode editar por produto no Excel
EDITAVEIS_MKT = {'MARGEM_ALVO', 'SH_PRECO', 'ML_PRECO', 'ML_FRETE_BASE'}
# Colunas em negrito (resultados de cada etapa)
DESTAQUE_MKT = {'C_ENT', 'C_BASE', 'P_BASE', 'B_LUCRO', 'SH_PRECO', 'SH_LUCRO', 'SH_MARGEM',
                'ML_PRECO', 'ML_LUCRO', 'ML_MARGEM', 'MELHOR'}


def valores_variaveis(M):
    """Valores das variáveis do topo das colunas."""
    return {**M, 'desc_rep': REPUTACAO_DESC[M['reputacao']]}


def _brl(v):
    return f"R$ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')


def _pct(v, casas=1):
    return f"{v*100:.{casas}f}%".replace('.', ',')


def _fmt_html(v, fmt):
    if fmt == FMT_BRL:
        return _brl(v)
    if fmt == FMT_PCT:
        return _pct(v, 2)
    if fmt == FMT_MULT:
        return f"× {v:.2f}".replace('.', ',')
    return html.escape(str(v))


def _valores_linha(row, m):
    """Valor de cada coluna para a prévia (mesmos números que o Excel calcula)."""
    sh, ml, b = m['shopee'], m['ml'], m['base']
    return {
        'NF': row['nf'], 'DESC': row['desc'][:55], 'REF': row['ref'], 'SKU': row['sku'],
        'QTD': int(row['qtd']), 'CST': row['cst'],
        'NF_U': row['nf_u'], 'C_REAL': m['c_real'], 'ST_U': row['st_u'], 'ANT_U': row['ant_u'],
        'IPI_U': row['ipi_u'], 'FRETE': m['frete'],
        'C_ENT': m['c_ent'], 'EMB': m['emb'], 'C_BASE': m['custo_base'],
        'MARGEM_ALVO': m['margem_alvo'], 'P_BASE': m['p_base'],
        'B_SIMPLES': b['simples'], 'B_DESP': b['desp'], 'B_LUCRO': b['lucro'],
        'SH_PRECO': sh['preco'], 'SH_COM_PCT': sh['com_pct'], 'SH_COM': sh['comissao'],
        'SH_TARIFA': sh['tarifa'], 'SH_SIMPLES': sh['simples'], 'SH_DESP': sh['desp'],
        'SH_LUCRO': sh['lucro'], 'SH_MARGEM': sh['margem'],
        'ML_PRECO': ml['preco'], 'ML_COM_PCT': ml['com_pct'], 'ML_COM': ml['comissao'],
        'ML_TARIFA': ml['tarifa'], 'ML_FRETE_BASE': m['frete_ml_base'], 'ML_FRETE': ml['frete'],
        'ML_SIMPLES': ml['simples'], 'ML_DESP': ml['desp'], 'ML_LUCRO': ml['lucro'], 'ML_MARGEM': ml['margem'],
        'MELHOR': m['melhor'],
    }


# ─── Prévia HTML ─────────────────────────────────────────────────────────────
def gerar_tabela_marketplace_html(rows, metricas, M):
    variaveis = valores_variaveis(M)
    out = '<table class="table table-sm table-bordered table-hover mkt-table"><thead>'

    # 1ª linha: grupos
    out += '<tr>'
    for g, rotulo, cor in GRUPOS_MKT:
        span = sum(1 for c in COLUNAS_MKT if c[2] == g)
        out += (f'<th colspan="{span}" style="background:#{cor};text-align:center;'
                f'font-size:0.7rem;letter-spacing:1px">{rotulo}</th>')
    out += '</tr>'

    # 2ª linha: cabeçalhos
    out += '<tr>' + ''.join(f'<th style="background:#{_COR_GRUPO[c[2]]}">{html.escape(c[1])}</th>'
                            for c in COLUNAS_MKT) + '</tr>'

    # 3ª linha: variáveis usadas em cada coluna
    out += '<tr class="mkt-vars">'
    for chave, _, grupo, _, var in COLUNAS_MKT:
        if var is None:
            txt = 'VARIÁVEIS →' if chave == 'DESC' else ''
            out += f'<th class="mkt-var-vazia">{txt}</th>'
            continue
        tipo = var[0]
        if tipo == 'param':
            txt, cls = _fmt_html(variaveis[var[1]], var[2]), 'mkt-var'
        elif tipo == 'link':
            origem = next(c for c in COLUNAS_MKT if c[0] == var[1])[4]
            txt, cls = _fmt_html(variaveis[origem[1]], origem[2]), 'mkt-var mkt-var-link'
        else:
            txt, cls = var[1], 'mkt-var mkt-var-link'
        out += f'<th class="{cls}">{txt}</th>'
    out += '</tr></thead><tbody>'

    for row, m in zip(rows, metricas):
        vals = _valores_linha(row, m)
        cor_linha = COR_ST if row['tem_st'] else (COR_ANT if row['tem_ant'] else None)
        out += '<tr>'
        for chave, _, grupo, fmt, _ in COLUNAS_MKT:
            v = vals[chave]
            if grupo in ('ID', 'FIM'):
                cor = cor_linha
            else:
                cor = _COR_GRUPO[grupo]
            estilo = f'background:#{cor};' if cor else ''
            if chave in DESTAQUE_MKT:
                estilo += 'font-weight:600;'
            if chave in ('SH_MARGEM', 'ML_MARGEM'):
                estilo += f'color:{"#15803d" if v >= m["margem_alvo"] - 1e-3 else "#dc2626"};'
                txt = _pct(v)
            elif chave in ('ST_U', 'ANT_U', 'IPI_U', 'ML_FRETE', 'SH_TARIFA', 'ML_TARIFA') and v < 0.005:
                txt = '-'
            elif fmt is None:
                txt = html.escape(str(v))
            else:
                txt = _fmt_html(v, fmt)
            out += f'<td style="{estilo}">{txt}</td>'
        out += '</tr>'
    return out + '</tbody></table>'


# ─── Dashboard HTML ──────────────────────────────────────────────────────────
def gerar_dashboard_marketplace_html(rows, metricas, M, num_nf=None):
    total  = len(rows)
    qtd    = [r['qtd'] for r in rows]
    inv    = sum(m['c_ent'] * q for m, q in zip(metricas, qtd))
    lucro_sh = sum(m['shopee']['lucro'] * q for m, q in zip(metricas, qtd))
    lucro_ml = sum(m['ml']['lucro'] * q for m, q in zip(metricas, qtd))
    fat_sh   = sum(m['shopee']['preco'] * q for m, q in zip(metricas, qtd))
    fat_ml   = sum(m['ml']['preco'] * q for m, q in zip(metricas, qtd))
    n_sh = sum(1 for m in metricas if m['melhor'] == MELHOR_SHOPEE)
    n_ml = total - n_sh
    com_frete = sum(1 for m in metricas if m['ml']['frete'] > 0)

    def card(titulo, valor, sub='', cor='#1e293b', fundo='#fff'):
        return (f'<div class="mkt-card" style="background:{fundo}">'
                f'<div class="mkt-card-t">{titulo}</div>'
                f'<div class="mkt-card-v" style="color:{cor}">{valor}</div>'
                f'<div class="mkt-card-s">{sub}</div></div>')

    rep = M['reputacao'].capitalize()
    out = '''
<style>
.mkt-dash{margin-top:24px}
.mkt-dash h6{font-size:.75rem;font-weight:700;color:#94a3b8;text-transform:uppercase;letter-spacing:1px;margin:22px 0 10px}
.mkt-cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:12px}
.mkt-card{border:1px solid #e2e8f0;border-radius:12px;padding:14px 16px}
.mkt-card-t{font-size:.72rem;font-weight:600;color:#64748b;text-transform:uppercase;letter-spacing:.5px}
.mkt-card-v{font-size:1.15rem;white-space:nowrap;font-weight:700;margin-top:4px;font-variant-numeric:tabular-nums}
.mkt-card-s{font-size:.75rem;color:#94a3b8;margin-top:2px}
.mkt-bars{display:flex;flex-direction:column;gap:10px}
.mkt-bar-row{display:grid;grid-template-columns:minmax(120px,260px) 1fr;gap:12px;align-items:center;font-size:.78rem}
.mkt-bar-nome{white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:#334155}
.mkt-bar-trilha{display:flex;flex-direction:column;gap:3px}
.mkt-bar{height:14px;border-radius:4px;display:flex;align-items:center;padding-left:6px;font-size:.68rem;font-weight:600;color:#1e293b;white-space:nowrap;min-width:2px}
.mkt-legenda{display:flex;gap:16px;font-size:.75rem;color:#64748b;margin-bottom:10px}
.mkt-legenda span::before{content:"";display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:5px;vertical-align:middle;background:var(--c)}
.mkt-alerta{background:#fffbeb;border-left:3px solid #f59e0b;padding:10px 14px;border-radius:0 8px 8px 0;font-size:.8rem;color:#92400e;margin-top:8px}
</style>
<div class="mkt-dash">'''
    titulo_nf = f' — NF {html.escape(str(num_nf))}' if num_nf else ''
    out += f'<h6>Resumo dos produtos selecionados{titulo_nf}</h6><div class="mkt-cards">'
    out += card('Produtos', total, f'{int(sum(qtd))} unidades')
    out += card('Custo de entrada', _brl(inv), 'total das unidades selecionadas')
    out += card('Lucro potencial Shopee', _brl(lucro_sh),
                f'faturamento {_brl(fat_sh)}', '#c2410c', f'#{COR_SHOPEE}40')
    out += card('Lucro potencial ML', _brl(lucro_ml),
                f'faturamento {_brl(fat_ml)}', '#92400e', f'#{COR_ML}60')
    out += card('Melhor canal', f'{n_sh} Shopee · {n_ml} ML', 'por lucro unitário')
    out += card('Margem alvo', _pct(M['margem']), f'reputação ML: {rep}')
    out += '</div>'

    # Barras: lucro unitário por produto em cada canal
    ordem = sorted(zip(rows, metricas),
                   key=lambda rm: max(rm[1]['shopee']['lucro'], rm[1]['ml']['lucro']),
                   reverse=True)
    maior = max([max(m['shopee']['lucro'], m['ml']['lucro']) for _, m in ordem] + [0.01])
    out += '<h6>Lucro por unidade — Shopee × Mercado Livre</h6>'
    out += (f'<div class="mkt-legenda"><span style="--c:#fdba74">Shopee</span>'
            f'<span style="--c:#fde047">Mercado Livre</span></div><div class="mkt-bars">')
    for row, m in ordem:
        def barra(v, cor):
            w = max(0.0, v) / maior * 100
            return (f'<div class="mkt-bar" style="width:{w:.1f}%;background:{cor}">'
                    f'{_brl(v)}</div>')
        out += (f'<div class="mkt-bar-row"><div class="mkt-bar-nome" title="{html.escape(row["desc"])}">'
                f'{html.escape(row["desc"][:40])}</div><div class="mkt-bar-trilha">'
                f'{barra(m["shopee"]["lucro"], "#fdba74")}{barra(m["ml"]["lucro"], "#fde047")}'
                f'</div></div>')
    out += '</div>'

    # Alertas
    alertas = []
    inviaveis = [r['desc'][:50] for r, m in zip(rows, metricas) if m['inviavel']]
    if inviaveis:
        alertas.append('Margem alvo inatingível (taxas + Simples + despesa + margem ≥ 100%): '
                       + ', '.join(html.escape(d) for d in inviaveis))
    if com_frete:
        alertas.append(f'{com_frete} produto(s) ficam a partir de R$ 79 no Mercado Livre e pagam frete '
                       f'estimado de {_brl(frete_ml_item(M))} (base {_brl(M["frete_ml"])}, '
                       f'reputação {rep}). Ajuste o frete por produto na planilha se souber o valor real.')
    if alertas:
        out += '<h6>Atenção</h6>' + ''.join(f'<div class="mkt-alerta">{a}</div>' for a in alertas)

    return out + '</div>'


# ─── Excel ───────────────────────────────────────────────────────────────────
def _fill(hex_color):
    return PatternFill('solid', start_color=hex_color)


_BORDA = Border(*(Side('thin'),) * 4)
_CTR   = Alignment(horizontal='center', vertical='center')
F_PARAM = _fill('F2F2F2')
F_EDIT  = _fill(COR_EDIT)
F_LINK  = _fill(COR_LINK)
F_ST    = _fill(COR_ST)
F_ANT   = _fill(COR_ANT)

# Linhas da aba Marketplace
LIN_GRUPO, LIN_HEADER, LIN_VAR, LIN_DADOS = 1, 2, 3, 4


def _cel(ws, r, c, value=None, fill=None, fmt=None, bold=False, italic=False):
    cell = ws.cell(row=r, column=c, value=value)
    cell.border = _BORDA
    cell.alignment = _CTR
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if bold or italic:
        cell.font = Font(bold=bold, italic=italic)
    return cell


def _aba_faixas(wb, M):
    """Aba com as tabelas de faixas (editáveis) usadas nos VLOOKUPs."""
    ws = wb.create_sheet('Faixas')
    f_sh, f_ml = _fill(COR_SHOPEE), _fill(COR_ML)

    for c, h in enumerate(['SHOPEE — A PARTIR DE', 'TARIFA FIXA', 'COMISSÃO'], 1):
        _cel(ws, 1, c, h, f_sh, bold=True)
    for i, (lo, fixo, pct) in enumerate(SHOPEE_FAIXAS, 2):
        _cel(ws, i, 1, lo, fmt=FMT_BRL)
        _cel(ws, i, 2, fixo, F_EDIT, FMT_BRL)
        _cel(ws, i, 3, pct, F_EDIT, FMT_PCT)

    for c, h in enumerate(['ML — A PARTIR DE', 'TARIFA FIXA', '% EXTRA (+ comissão)'], 5):
        _cel(ws, 1, c, h, f_ml, bold=True)
    for i, (lo, fixo, pct) in enumerate(ML_FAIXAS, 2):
        _cel(ws, i, 5, lo, fmt=FMT_BRL)
        _cel(ws, i, 6, fixo, F_EDIT, FMT_BRL)
        _cel(ws, i, 7, pct, F_EDIT, FMT_PCT)

    _cel(ws, 1, 9, 'ML: VENDEDOR PAGA FRETE A PARTIR DE', f_ml, bold=True)
    _cel(ws, 2, 9, ML_LIMITE_FRETE, F_EDIT, FMT_BRL)
    _cel(ws, 4, 9, 'REPUTAÇÃO ML', f_ml, bold=True)
    _cel(ws, 4, 10, 'DESCONTO NO FRETE', f_ml, bold=True)
    for i, (rep, desc) in enumerate(REPUTACAO_DESC.items(), 5):
        atual = rep == M['reputacao']
        _cel(ws, i, 9, rep.capitalize() + ('  ← atual' if atual else ''), bold=atual)
        _cel(ws, i, 10, desc, fmt=FMT_PCT, bold=atual)

    for col, w in zip('ABCDEFGHIJ', (22, 14, 12, 3, 20, 14, 20, 3, 36, 20)):
        ws.column_dimensions[col].width = w
    return {
        'shopee': f"Faixas!$A$2:$C${1 + len(SHOPEE_FAIXAS)}",
        'ml':     f"Faixas!$E$2:$G${1 + len(ML_FAIXAS)}",
        'lim':    "Faixas!$I$2",
    }


def _formulas_linha(c, v, t):
    """
    Fórmulas de uma linha. `c` = célula da coluna nesta linha, `v` = célula da
    variável no topo da coluna (linha 3, absoluta), `t` = tabelas de faixas.
    """
    simp, desp, marg = v['B_SIMPLES'], v['B_DESP'], c['MARGEM_ALVO']
    # Preço base: margem alvo só com Simples + despesa (mesma conta de _resolver)
    den = f'(1-{simp}-{desp}-{marg})'

    sh, ml = c['SH_PRECO'], c['ML_PRECO']
    return {
        'C_REAL':  f'=ROUND({c["NF_U"]}*{v["C_REAL"]},2)',
        'FRETE':   f'=ROUND({c["C_REAL"]}*{v["FRETE"]},2)',
        'C_ENT':   (f'=ROUND({c["C_REAL"]}+{c["ST_U"]}+{c["ANT_U"]}+{c["IPI_U"]}'
                    f'+{c["FRETE"]},2)'),
        'EMB':     f'={v["EMB"]}',
        'C_BASE':  f'=ROUND({c["C_ENT"]}+{c["EMB"]},2)',
        'MARGEM_ALVO': f'={v["MARGEM_ALVO"]}',
        # Arredonda para cima até X,X9 — mesmo que arredondar_x9_acima()
        'P_BASE':  (f'=IF({den}<=0,0,ROUND(ROUNDUP(ROUND(({c["C_BASE"]}/{den}-0.09)*10,6),0)'
                    f'/10+0.09,2))'),
        'B_SIMPLES': f'=ROUND({c["P_BASE"]}*{simp},2)',
        'B_DESP':    f'=ROUND({c["P_BASE"]}*{desp},2)',
        'B_LUCRO': f'=ROUND({c["P_BASE"]}-{c["C_BASE"]}-{c["B_SIMPLES"]}-{c["B_DESP"]},2)',

        'SH_COM_PCT': f'=VLOOKUP({sh},{t["shopee"]},3,TRUE)',
        'SH_COM':     f'=ROUND({sh}*{c["SH_COM_PCT"]},2)',
        'SH_TARIFA':  f'=VLOOKUP({sh},{t["shopee"]},2,TRUE)',
        'SH_SIMPLES': f'=ROUND({sh}*{simp},2)',
        'SH_DESP':    f'=ROUND({sh}*{desp},2)',
        'SH_LUCRO':   (f'=ROUND({sh}-{c["C_BASE"]}-{c["SH_COM"]}-{c["SH_TARIFA"]}'
                       f'-{c["SH_SIMPLES"]}-{c["SH_DESP"]},2)'),
        'SH_MARGEM':  f'=IF({sh}>0,ROUND({c["SH_LUCRO"]}/{sh},4),0)',

        'ML_COM_PCT': f'={v["ML_COM_PCT"]}+VLOOKUP({ml},{t["ml"]},3,TRUE)',
        'ML_COM':     f'=ROUND({ml}*{c["ML_COM_PCT"]},2)',
        'ML_TARIFA':  f'=IF({ml}>={t["lim"]},0,VLOOKUP({ml},{t["ml"]},2,TRUE))',
        'ML_FRETE_BASE': f'={v["ML_FRETE_BASE"]}',
        'ML_FRETE':   f'=IF({ml}>={t["lim"]},ROUND({c["ML_FRETE_BASE"]}*(1-{v["ML_FRETE"]}),2),0)',
        'ML_SIMPLES': f'=ROUND({ml}*{simp},2)',
        'ML_DESP':    f'=ROUND({ml}*{desp},2)',
        'ML_LUCRO':   (f'=ROUND({ml}-{c["C_BASE"]}-{c["ML_COM"]}-{c["ML_TARIFA"]}-{c["ML_FRETE"]}'
                       f'-{c["ML_SIMPLES"]}-{c["ML_DESP"]},2)'),
        'ML_MARGEM':  f'=IF({ml}>0,ROUND({c["ML_LUCRO"]}/{ml},4),0)',
        'MELHOR':     f'=IF({c["SH_LUCRO"]}>={c["ML_LUCRO"]},"{MELHOR_SHOPEE}","{MELHOR_ML}")',
    }


def salvar_excel_marketplace(rows, metricas, M, num_nf, path):
    L = get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Marketplace'
    tabelas = _aba_faixas(wb, M)
    variaveis = valores_variaveis(M)

    # Linha 1: grupos
    for g, rotulo, cor in GRUPOS_MKT:
        cols = [COL_MKT[c[0]] for c in COLUNAS_MKT if c[2] == g]
        if len(cols) > 1:
            ws.merge_cells(start_row=LIN_GRUPO, start_column=cols[0],
                           end_row=LIN_GRUPO, end_column=cols[-1])
        _cel(ws, LIN_GRUPO, cols[0], rotulo, _fill(cor), bold=True)

    # Linha 2: cabeçalhos
    for chave, header, grupo, _, _ in COLUNAS_MKT:
        _cel(ws, LIN_HEADER, COL_MKT[chave], header, _fill(_COR_GRUPO[grupo]), bold=True)

    # Linha 3: variáveis no topo de cada coluna (editáveis em verde)
    v = {chave: f'${L(COL_MKT[chave])}${LIN_VAR}' for chave in COL_MKT}
    for chave, _, _, _, var in COLUNAS_MKT:
        col = COL_MKT[chave]
        if var is None:
            _cel(ws, LIN_VAR, col, 'VARIÁVEIS →' if chave == 'DESC' else None, F_PARAM, bold=True)
        elif var[0] == 'param':
            _cel(ws, LIN_VAR, col, variaveis[var[1]], F_EDIT, var[2], bold=True)
        elif var[0] == 'link':
            origem = next(c for c in COLUNAS_MKT if c[0] == var[1])[4]
            _cel(ws, LIN_VAR, col, f'={v[var[1]]}', F_LINK, origem[2], italic=True)
        else:
            _cel(ws, LIN_VAR, col, var[1], F_LINK, italic=True)

    # Dados
    for i, (row, m) in enumerate(zip(rows, metricas)):
        r = LIN_DADOS + i
        c = {k: f'{L(col)}{r}' for k, col in COL_MKT.items()}
        formulas = _formulas_linha(c, v, tabelas)
        fixos = {
            'NF': row['nf'], 'DESC': row['desc'], 'REF': row['ref'], 'SKU': row['sku'],
            'QTD': row['qtd'], 'CST': str(row['cst']),
            'NF_U': row['nf_u'], 'ST_U': row['st_u'], 'ANT_U': row['ant_u'], 'IPI_U': row['ipi_u'],
            # Preços sugeridos vêm do Python (a faixa depende do próprio preço) — editáveis
            'SH_PRECO': m['shopee']['preco'], 'ML_PRECO': m['ml']['preco'],
        }
        fill_linha = F_ST if row['tem_st'] else (F_ANT if row['tem_ant'] else None)
        for chave, _, grupo, fmt, _ in COLUNAS_MKT:
            val = formulas.get(chave, fixos.get(chave))
            if chave in EDITAVEIS_MKT:
                fill = F_EDIT
            elif grupo in ('ID', 'FIM'):
                fill = fill_linha
            else:
                fill = _fill(_COR_GRUPO[grupo])
            _cel(ws, r, COL_MKT[chave], val, fill, fmt, chave in DESTAQUE_MKT)
        ws.cell(r, COL_MKT['DESC']).alignment = Alignment(vertical='center')

    larguras = {'DESC': 45, 'NF': 8, 'QTD': 7, 'CST': 7, 'MELHOR': 16}
    for chave, col in COL_MKT.items():
        ws.column_dimensions[L(col)].width = larguras.get(chave, 13)
    ws.row_dimensions[LIN_HEADER].height = 32
    for cell in ws[LIN_HEADER]:
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(LIN_DADOS, COL_MKT['REF'])

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
